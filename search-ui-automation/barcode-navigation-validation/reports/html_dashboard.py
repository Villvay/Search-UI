"""HTML dashboard for barcode navigation validation."""

from __future__ import annotations

import json
from pathlib import Path
from typing import Any

from constants import STATUS_FAIL, STATUS_PASS

STATUS_NOT_TESTED = "NOT_TESTED"
STATUS_NOT_MAPPED = "NOT_MAPPED"


def _trim(value: Any, limit: int = 80) -> str:
    text = str(value or "")
    return text if len(text) <= limit else text[: limit - 1] + "…"


def _normalize_ident(value: Any) -> str:
    ident = str(value or "Unknown")
    return "materialNumber" if ident == "SKU" else ident


def _split_alias(value: str) -> list[str]:
    parts: list[str] = []
    for chunk in str(value or "").replace(";", ",").split(","):
        item = chunk.strip()
        if item:
            parts.append(item)
    return parts


def load_product_families(product_dataset_path: Path | None) -> list[dict[str, Any]]:
    """Load product rows keyed by materialNumber with related identifiers."""
    if product_dataset_path is None or not product_dataset_path.is_file():
        return []

    from openpyxl import load_workbook

    families: list[dict[str, Any]] = []
    workbook = load_workbook(product_dataset_path, read_only=True, data_only=True)
    try:
        sheet = (
            workbook["Products"]
            if "Products" in workbook.sheetnames
            else workbook.active
        )
        rows = sheet.iter_rows(values_only=True)
        header = next(rows, None)
        if not header:
            return []
        headers = [str(cell or "").strip() for cell in header]
        idx = {name: i for i, name in enumerate(headers)}

        def _cell(row: tuple, name: str) -> str:
            if name not in idx or idx[name] >= len(row):
                return ""
            return "" if row[idx[name]] is None else str(row[idx[name]]).strip()

        for row in rows:
            if not row:
                continue
            material = _cell(row, "materialNumber")
            if not material:
                continue
            families.append(
                {
                    "materialNumber": material,
                    "MFRPartNo": _cell(row, "MFRPartNo"),
                    "UPCCode": _cell(row, "UPCCode"),
                    "Alias": _split_alias(_cell(row, "Alias")),
                    "Product": _cell(row, "Product"),
                }
            )
    finally:
        workbook.close()
    return families


def _status_lookup(rows: list[dict[str, Any]]) -> dict[str, dict[str, Any]]:
    """Map barcode value -> best result row (prefer FAIL over PASS if dupes)."""
    lookup: dict[str, dict[str, Any]] = {}
    for row in rows:
        barcode = str(row.get("barcode") or "").strip()
        if not barcode:
            continue
        current = lookup.get(barcode)
        if current is None:
            lookup[barcode] = row
            continue
        # Prefer FAIL if the same value appears with conflicting outcomes.
        if current.get("status") != STATUS_FAIL and row.get("status") == STATUS_FAIL:
            lookup[barcode] = row
    return lookup


def _field_status(
    value: str, status_by_barcode: dict[str, dict[str, Any]]
) -> dict[str, str]:
    text = (value or "").strip()
    if not text:
        return {
            "value": "",
            "status": STATUS_NOT_MAPPED,
            "reason": "",
            "identifier_type": "",
        }
    hit = status_by_barcode.get(text)
    if not hit:
        return {
            "value": text,
            "status": STATUS_NOT_TESTED,
            "reason": "",
            "identifier_type": "",
        }
    return {
        "value": text,
        "status": str(hit.get("status") or STATUS_NOT_TESTED),
        "reason": _trim(hit.get("reason", ""), 60),
        "identifier_type": _normalize_ident(hit.get("identifier_type")),
    }


def build_related_families(
    rows: list[dict[str, Any]],
    product_families: list[dict[str, Any]],
) -> list[dict[str, Any]]:
    """Attach PASS/FAIL for material / MFR / UPC / Alias on each product family."""
    status_by_barcode = _status_lookup(rows)
    related: list[dict[str, Any]] = []

    for family in product_families:
        material = _field_status(family.get("materialNumber", ""), status_by_barcode)
        mfr = _field_status(family.get("MFRPartNo", ""), status_by_barcode)
        upc = _field_status(family.get("UPCCode", ""), status_by_barcode)
        aliases = [
            _field_status(alias, status_by_barcode)
            for alias in (family.get("Alias") or [])
        ]

        tested_statuses = [
            item["status"]
            for item in [material, mfr, upc, *aliases]
            if item["status"] in {STATUS_PASS, STATUS_FAIL}
        ]
        if not tested_statuses:
            continue

        any_fail = STATUS_FAIL in tested_statuses
        any_pass = STATUS_PASS in tested_statuses
        related.append(
            {
                "materialNumber": material["value"],
                "product_name": _trim(family.get("Product", ""), 80),
                "material": material,
                "mfr": mfr,
                "upc": upc,
                "aliases": aliases,
                "any_failed": any_fail,
                "any_passed": any_pass,
                "mixed": any_fail and any_pass,
                "all_failed": bool(tested_statuses) and all(
                    s == STATUS_FAIL for s in tested_statuses
                ),
                "all_passed": bool(tested_statuses) and all(
                    s == STATUS_PASS for s in tested_statuses
                ),
            }
        )

    related.sort(
        key=lambda item: (
            0 if item["mixed"] else 1 if item["any_failed"] else 2,
            item["materialNumber"],
        )
    )
    return related


def _enrich_rows_with_related(
    slim_rows: list[dict[str, Any]],
    product_families: list[dict[str, Any]],
    status_by_barcode: dict[str, dict[str, Any]],
) -> None:
    """Add related-field status columns onto each result row."""
    by_material: dict[str, dict[str, Any]] = {}
    by_any_value: dict[str, dict[str, Any]] = {}
    for family in product_families:
        material = str(family.get("materialNumber") or "").strip()
        if not material:
            continue
        by_material[material] = family
        for key in ("materialNumber", "MFRPartNo", "UPCCode"):
            value = str(family.get(key) or "").strip()
            if value:
                by_any_value[value] = family
        for alias in family.get("Alias") or []:
            text = str(alias or "").strip()
            if text:
                by_any_value[text] = family

    for row in slim_rows:
        barcode = str(row.get("barcode") or "").strip()
        family = by_any_value.get(barcode) or by_material.get(barcode)
        if not family:
            row["related_materialNumber"] = ""
            row["related_material_status"] = STATUS_NOT_MAPPED
            row["related_mfr"] = ""
            row["related_mfr_status"] = STATUS_NOT_MAPPED
            row["related_upc"] = ""
            row["related_upc_status"] = STATUS_NOT_MAPPED
            row["related_alias"] = ""
            row["related_alias_status"] = STATUS_NOT_MAPPED
            row["related_summary"] = ""
            continue

        material = _field_status(family.get("materialNumber", ""), status_by_barcode)
        mfr = _field_status(family.get("MFRPartNo", ""), status_by_barcode)
        upc = _field_status(family.get("UPCCode", ""), status_by_barcode)
        aliases = [
            _field_status(alias, status_by_barcode)
            for alias in (family.get("Alias") or [])
        ]
        alias_text = "; ".join(
            f"{item['value']} ({item['status']})" for item in aliases if item["value"]
        )
        alias_status = ""
        if not aliases:
            alias_status = STATUS_NOT_MAPPED
        elif any(item["status"] == STATUS_FAIL for item in aliases):
            alias_status = STATUS_FAIL
        elif any(item["status"] == STATUS_PASS for item in aliases):
            alias_status = STATUS_PASS
        elif aliases:
            alias_status = aliases[0]["status"]

        row["related_materialNumber"] = material["value"]
        row["related_material_status"] = material["status"]
        row["related_mfr"] = mfr["value"]
        row["related_mfr_status"] = mfr["status"]
        row["related_upc"] = upc["value"]
        row["related_upc_status"] = upc["status"]
        row["related_alias"] = alias_text
        row["related_alias_status"] = alias_status
        row["related_summary"] = (
            f"MAT:{material['status']} | MFR:{mfr['status']} | "
            f"UPC:{upc['status']} | ALIAS:{alias_status or STATUS_NOT_MAPPED}"
        )


def build_dashboard(
    rows: list[dict[str, Any]],
    summary: dict[str, Any],
    *,
    product_dataset_path: Path | None = None,
) -> str:
    slim_rows: list[dict[str, Any]] = []
    for r in rows:
        slim_rows.append(
            {
                "barcode": r.get("barcode", ""),
                "identifier_type": _normalize_ident(r.get("identifier_type")),
                "status": r.get("status", ""),
                "reason": _trim(r.get("reason", ""), 60),
                "http_status": r.get("http_status", ""),
                "product_id": r.get("product_id", ""),
                "product_name": _trim(r.get("product_name", ""), 80),
                "execution_time_ms": r.get("execution_time_ms", ""),
            }
        )

    product_families = load_product_families(product_dataset_path)
    status_by_barcode = _status_lookup(slim_rows)
    _enrich_rows_with_related(slim_rows, product_families, status_by_barcode)
    related_families = build_related_families(slim_rows, product_families)
    mixed_families = [f for f in related_families if f.get("mixed")]
    failed_families = [f for f in related_families if f.get("any_failed")]

    passed = [r for r in slim_rows if r.get("status") == STATUS_PASS]
    failed = [r for r in slim_rows if r.get("status") == STATUS_FAIL]
    rows_json = json.dumps(slim_rows, ensure_ascii=False)
    passed_json = json.dumps(passed, ensure_ascii=False)
    failed_json = json.dumps(failed, ensure_ascii=False)
    related_json = json.dumps(related_families, ensure_ascii=False)
    mixed_json = json.dumps(mixed_families, ensure_ascii=False)
    failed_related_json = json.dumps(failed_families, ensure_ascii=False)
    summary_json = json.dumps(summary, ensure_ascii=False)

    type_stats: dict[str, dict[str, int]] = {}
    for row in slim_rows:
        key = str(row.get("identifier_type") or "Unknown")
        bucket = type_stats.setdefault(key, {"total": 0, "passed": 0, "failed": 0})
        bucket["total"] += 1
        if row.get("status") == STATUS_PASS:
            bucket["passed"] += 1
        elif row.get("status") == STATUS_FAIL:
            bucket["failed"] += 1
    type_stats_json = json.dumps(type_stats, ensure_ascii=False)

    related_stats = {
        "families_with_tests": len(related_families),
        "mixed_families": len(mixed_families),
        "families_with_any_fail": len(failed_families),
        "product_dataset": str(product_dataset_path) if product_dataset_path else "",
    }
    related_stats_json = json.dumps(related_stats, ensure_ascii=False)

    return f"""<!doctype html>
<html lang="en">
<head>
  <meta charset="utf-8" />
  <meta name="viewport" content="width=device-width, initial-scale=1" />
  <title>Barcode Navigation Validation Dashboard</title>
  <style>
    body {{ font-family: Arial, sans-serif; margin: 20px; color: #111; }}
    h1 {{ margin-bottom: 8px; }}
    h2 {{ margin: 22px 0 8px; font-size: 18px; }}
    .cards {{ display: flex; gap: 12px; flex-wrap: wrap; margin: 14px 0; }}
    .card {{ border: 1px solid #ddd; border-radius: 8px; padding: 10px 14px; min-width: 180px; }}
    .card .sub {{ margin-top: 6px; font-size: 13px; color: #444; line-height: 1.45; }}
    .pass-text {{ color: #047857; font-weight: 700; }}
    .fail-text {{ color: #b91c1c; font-weight: 700; }}
    .neutral-text {{ color: #6b7280; font-weight: 700; }}
    .tabs button {{ margin-right: 8px; margin-bottom: 8px; padding: 8px 10px; border: 1px solid #bbb; background: #f5f5f5; border-radius: 6px; cursor: pointer; }}
    .tabs button.active {{ background: #111; color: #fff; border-color: #111; }}
    .toolbar {{ margin: 10px 0 6px; display: flex; gap: 10px; flex-wrap: wrap; align-items: center; }}
    .toolbar input, .toolbar select {{ padding: 7px 9px; border: 1px solid #ccc; border-radius: 6px; }}
    .toolbar input {{ min-width: 240px; }}
    table {{ width: 100%; border-collapse: collapse; margin-top: 10px; }}
    th, td {{ border: 1px solid #ddd; padding: 7px 8px; text-align: left; font-size: 13px; vertical-align: top; }}
    th {{ background: #f6f6f6; position: sticky; top: 0; z-index: 1; }}
    .pass {{ color: #047857; font-weight: 700; }}
    .fail {{ color: #b91c1c; font-weight: 700; }}
    .not-tested {{ color: #6b7280; font-weight: 600; }}
    .muted {{ color: #555; font-size: 13px; }}
    .count {{ color: #666; font-size: 13px; }}
    .type {{ display: inline-block; padding: 2px 8px; border-radius: 999px; background: #eef2ff; font-size: 12px; }}
    .pill {{ display: inline-block; padding: 2px 8px; border-radius: 999px; font-size: 12px; font-weight: 700; }}
    .pill-pass {{ background: #d1fae5; color: #047857; }}
    .pill-fail {{ background: #fee2e2; color: #b91c1c; }}
    .pill-other {{ background: #f3f4f6; color: #4b5563; }}
    .panel {{ display: none; }}
    .panel.active {{ display: block; }}
    .alias-list {{ margin: 0; padding-left: 16px; }}
    tr.mixed {{ background: #fff7ed; }}
  </style>
</head>
<body>
  <h1>Barcode Navigation Validation <span id="versionBadge" class="type"></span></h1>
  <div id="meta" class="muted"></div>
  <div class="cards" id="cards"></div>
  <div class="tabs">
    <button id="btnFailed" class="active" data-panel="results">Failed</button>
    <button id="btnPassed" data-panel="results">Passed</button>
    <button id="btnAll" data-panel="results">All</button>
    <button id="btnRelatedMixed" data-panel="related">Related: Mixed</button>
    <button id="btnRelatedFailed" data-panel="related">Related: Any Fail</button>
    <button id="btnRelatedAll" data-panel="related">Related: All Families</button>
  </div>

  <div id="panelResults" class="panel active">
    <div class="toolbar">
      <input id="filter" type="search" placeholder="Filter by barcode, type, reason, related fields..." />
      <select id="typeFilter">
        <option value="">All identifier types</option>
      </select>
      <span id="rowCount" class="count"></span>
    </div>
    <table>
      <thead>
        <tr>
          <th>barcode</th>
          <th>identifier_type</th>
          <th>status</th>
          <th>related_material</th>
          <th>related_mfr</th>
          <th>related_upc</th>
          <th>related_alias</th>
          <th>reason</th>
          <th>http_status</th>
          <th>product_id</th>
          <th>product_name</th>
          <th>execution_time_ms</th>
        </tr>
      </thead>
      <tbody id="tbody"></tbody>
    </table>
  </div>

  <div id="panelRelated" class="panel">
    <p class="muted">
      Each row is a product family from the product dataset. Statuses show whether
      materialNumber, MFRPartNo, UPCCode, and Alias values passed or failed in this run.
      <strong>Mixed</strong> means some related identifiers passed while others failed.
    </p>
    <div class="toolbar">
      <input id="relatedFilter" type="search" placeholder="Filter by material, MFR, UPC, alias..." />
      <span id="relatedCount" class="count"></span>
    </div>
    <table>
      <thead>
        <tr>
          <th>materialNumber</th>
          <th>material status</th>
          <th>MFRPartNo</th>
          <th>MFR status</th>
          <th>UPCCode</th>
          <th>UPC status</th>
          <th>Alias (status)</th>
          <th>outcome</th>
          <th>product_name</th>
        </tr>
      </thead>
      <tbody id="relatedBody"></tbody>
    </table>
  </div>

  <script>
    const allRows = {rows_json};
    const passedRows = {passed_json};
    const failedRows = {failed_json};
    const relatedAll = {related_json};
    const relatedMixed = {mixed_json};
    const relatedFailed = {failed_related_json};
    const summary = {summary_json};
    const typeStats = {type_stats_json};
    const relatedStats = {related_stats_json};
    let activeRows = failedRows;
    let activeRelated = relatedMixed;
    let activePanel = "results";

    const versionLabel = summary.version_label || summary.version || "";
    const versionKind = summary.dashboard_kind || summary.kind || "";
    const versionBadge = document.getElementById("versionBadge");
    versionBadge.textContent = versionLabel
      ? (versionKind ? (versionLabel + " · " + versionKind) : String(versionLabel))
      : "unversioned";

    const meta = document.getElementById("meta");
    meta.innerText =
      "Version: " + (versionLabel || "unversioned") +
      (versionKind ? (" (" + versionKind + ")") : "") +
      " | API: " + (summary.barcode_api_url || "") +
      " | Environment: " + (summary.environment || "") +
      " | Generated: " + (summary.generated_at || "") +
      " | Input: " + (summary.barcodes_csv || "");
    document.title = versionLabel
      ? ("Barcode Navigation Validation " + versionLabel)
      : "Barcode Navigation Validation Dashboard";

    const cards = document.getElementById("cards");
    function addCard(title, value, subHtml) {{
      const d = document.createElement("div");
      d.className = "card";
      d.innerHTML =
        "<div><strong>" + title + "</strong></div>" +
        "<div style='font-size:22px;margin-top:4px;'>" + (value ?? "") + "</div>" +
        (subHtml ? "<div class='sub'>" + subHtml + "</div>" : "");
      cards.appendChild(d);
    }}
    addCard("Total Tested", summary.total_tested);
    addCard("Passed", summary.passed);
    addCard("Failed", summary.failed);
    addCard("Execution (s)", summary.execution_time_s);
    addCard(
      "Related Families",
      relatedStats.families_with_tests || 0,
      "<span class='fail-text'>Mixed: " + (relatedStats.mixed_families || 0) + "</span><br>" +
      "<span class='neutral-text'>Any fail: " + (relatedStats.families_with_any_fail || 0) + "</span>"
    );

    const preferredOrder = ["materialNumber", "MFRPartNo", "UPCCode", "Alias"];
    const typeKeys = Object.keys(typeStats).sort((a, b) => {{
      const ai = preferredOrder.indexOf(a);
      const bi = preferredOrder.indexOf(b);
      if (ai === -1 && bi === -1) return a.localeCompare(b);
      if (ai === -1) return 1;
      if (bi === -1) return -1;
      return ai - bi;
    }});
    typeKeys.forEach(k => {{
      const s = typeStats[k] || {{}};
      addCard(
        k,
        s.total || 0,
        "<span class='pass-text'>Passed: " + (s.passed || 0) + "</span><br>" +
        "<span class='fail-text'>Failed: " + (s.failed || 0) + "</span>"
      );
    }});

    const typeFilter = document.getElementById("typeFilter");
    typeKeys.forEach(k => {{
      const s = typeStats[k] || {{}};
      const opt = document.createElement("option");
      opt.value = k;
      opt.textContent = k + " (P:" + (s.passed || 0) + " / F:" + (s.failed || 0) + ")";
      typeFilter.appendChild(opt);
    }});

    function escapeHtml(value) {{
      return String(value ?? "")
        .replaceAll("&", "&amp;")
        .replaceAll("<", "&lt;")
        .replaceAll(">", "&gt;")
        .replaceAll('"', "&quot;");
    }}

    function statusClass(status) {{
      if (status === "PASS") return "pass";
      if (status === "FAIL") return "fail";
      return "not-tested";
    }}

    function statusPill(status) {{
      const cls = status === "PASS" ? "pill-pass"
        : status === "FAIL" ? "pill-fail"
        : "pill-other";
      return "<span class='pill " + cls + "'>" + escapeHtml(status || "") + "</span>";
    }}

    function relatedCell(value, status) {{
      if (!value && (status === "NOT_MAPPED" || !status)) {{
        return "<span class='muted'>—</span>";
      }}
      return "<div>" + escapeHtml(value || "") + "</div>" + statusPill(status || "NOT_TESTED");
    }}

    function showPanel(name) {{
      activePanel = name;
      document.getElementById("panelResults").classList.toggle("active", name === "results");
      document.getElementById("panelRelated").classList.toggle("active", name === "related");
    }}

    function render(data) {{
      const q = (document.getElementById("filter").value || "").trim().toLowerCase();
      const type = typeFilter.value;
      const filtered = data.filter(r => {{
        if (type && r.identifier_type !== type) return false;
        if (!q) return true;
        const hay = [
          r.barcode, r.identifier_type, r.status, r.reason, r.http_status,
          r.product_id, r.product_name, r.related_summary,
          r.related_materialNumber, r.related_mfr, r.related_upc, r.related_alias
        ].join(" ").toLowerCase();
        return hay.includes(q);
      }});
      document.getElementById("rowCount").innerText = filtered.length + " row(s)";
      const tb = document.getElementById("tbody");
      tb.innerHTML = "";
      filtered.forEach(r => {{
        const tr = document.createElement("tr");
        const statusClassName = statusClass(r.status);
        tr.innerHTML = `
          <td>${{escapeHtml(r.barcode)}}</td>
          <td><span class="type">${{escapeHtml(r.identifier_type)}}</span></td>
          <td><span class="${{statusClassName}}">${{escapeHtml(r.status)}}</span></td>
          <td>${{relatedCell(r.related_materialNumber, r.related_material_status)}}</td>
          <td>${{relatedCell(r.related_mfr, r.related_mfr_status)}}</td>
          <td>${{relatedCell(r.related_upc, r.related_upc_status)}}</td>
          <td>${{relatedCell(r.related_alias, r.related_alias_status)}}</td>
          <td>${{escapeHtml(r.reason)}}</td>
          <td>${{escapeHtml(r.http_status)}}</td>
          <td>${{escapeHtml(r.product_id)}}</td>
          <td>${{escapeHtml(r.product_name)}}</td>
          <td>${{escapeHtml(r.execution_time_ms)}}</td>
        `;
        tb.appendChild(tr);
      }});
    }}

    function outcomeLabel(family) {{
      if (family.mixed) return "<span class='pill pill-fail'>MIXED</span>";
      if (family.all_failed) return "<span class='pill pill-fail'>ALL FAIL</span>";
      if (family.all_passed) return "<span class='pill pill-pass'>ALL PASS</span>";
      if (family.any_failed) return "<span class='pill pill-fail'>HAS FAIL</span>";
      return "<span class='pill pill-pass'>HAS PASS</span>";
    }}

    function renderRelated(data) {{
      const q = (document.getElementById("relatedFilter").value || "").trim().toLowerCase();
      const filtered = data.filter(f => {{
        if (!q) return true;
        const aliasHay = (f.aliases || []).map(a => (a.value || "") + " " + (a.status || "")).join(" ");
        const hay = [
          f.materialNumber, f.product_name,
          f.material && f.material.value, f.material && f.material.status,
          f.mfr && f.mfr.value, f.mfr && f.mfr.status,
          f.upc && f.upc.value, f.upc && f.upc.status,
          aliasHay
        ].join(" ").toLowerCase();
        return hay.includes(q);
      }});
      document.getElementById("relatedCount").innerText = filtered.length + " family(ies)";
      const tb = document.getElementById("relatedBody");
      tb.innerHTML = "";
      filtered.forEach(f => {{
        const tr = document.createElement("tr");
        if (f.mixed) tr.className = "mixed";
        const aliasHtml = (f.aliases && f.aliases.length)
          ? "<ul class='alias-list'>" + f.aliases.map(a =>
              "<li>" + escapeHtml(a.value) + " " + statusPill(a.status) + "</li>"
            ).join("") + "</ul>"
          : "<span class='muted'>—</span>";
        tr.innerHTML = `
          <td><strong>${{escapeHtml(f.materialNumber)}}</strong></td>
          <td>${{statusPill(f.material && f.material.status)}}</td>
          <td>${{escapeHtml(f.mfr && f.mfr.value || "") || "<span class='muted'>—</span>"}}</td>
          <td>${{statusPill(f.mfr && f.mfr.status)}}</td>
          <td>${{escapeHtml(f.upc && f.upc.value || "") || "<span class='muted'>—</span>"}}</td>
          <td>${{statusPill(f.upc && f.upc.status)}}</td>
          <td>${{aliasHtml}}</td>
          <td>${{outcomeLabel(f)}}</td>
          <td>${{escapeHtml(f.product_name || "")}}</td>
        `;
        tb.appendChild(tr);
      }});
    }}

    const btnFailed = document.getElementById("btnFailed");
    const btnPassed = document.getElementById("btnPassed");
    const btnAll = document.getElementById("btnAll");
    const btnRelatedMixed = document.getElementById("btnRelatedMixed");
    const btnRelatedFailed = document.getElementById("btnRelatedFailed");
    const btnRelatedAll = document.getElementById("btnRelatedAll");
    const resultButtons = [btnFailed, btnPassed, btnAll];
    const relatedButtons = [btnRelatedMixed, btnRelatedFailed, btnRelatedAll];

    function activate(btn, group) {{
      group.forEach(b => b.classList.remove("active"));
      btn.classList.add("active");
    }}

    btnFailed.onclick = () => {{
      relatedButtons.forEach(b => b.classList.remove("active"));
      activate(btnFailed, resultButtons);
      activeRows = failedRows;
      showPanel("results");
      render(activeRows);
    }};
    btnPassed.onclick = () => {{
      relatedButtons.forEach(b => b.classList.remove("active"));
      activate(btnPassed, resultButtons);
      activeRows = passedRows;
      showPanel("results");
      render(activeRows);
    }};
    btnAll.onclick = () => {{
      relatedButtons.forEach(b => b.classList.remove("active"));
      activate(btnAll, resultButtons);
      activeRows = allRows;
      showPanel("results");
      render(activeRows);
    }};
    btnRelatedMixed.onclick = () => {{
      resultButtons.forEach(b => b.classList.remove("active"));
      activate(btnRelatedMixed, relatedButtons);
      activeRelated = relatedMixed;
      showPanel("related");
      renderRelated(activeRelated);
    }};
    btnRelatedFailed.onclick = () => {{
      resultButtons.forEach(b => b.classList.remove("active"));
      activate(btnRelatedFailed, relatedButtons);
      activeRelated = relatedFailed;
      showPanel("related");
      renderRelated(activeRelated);
    }};
    btnRelatedAll.onclick = () => {{
      resultButtons.forEach(b => b.classList.remove("active"));
      activate(btnRelatedAll, relatedButtons);
      activeRelated = relatedAll;
      showPanel("related");
      renderRelated(activeRelated);
    }};

    document.getElementById("filter").addEventListener("input", () => render(activeRows));
    typeFilter.addEventListener("change", () => render(activeRows));
    document.getElementById("relatedFilter").addEventListener("input", () => renderRelated(activeRelated));
    render(activeRows);
  </script>
</body>
</html>
"""


def write_dashboard_html(
    path: Path,
    rows: list[dict[str, Any]],
    summary: dict[str, Any],
    *,
    product_dataset_path: Path | None = None,
) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(
        build_dashboard(
            rows,
            summary,
            product_dataset_path=product_dataset_path,
        ),
        encoding="utf-8",
    )
