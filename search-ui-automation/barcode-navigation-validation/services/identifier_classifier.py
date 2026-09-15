"""Classify barcode values against the product master dataset."""

from __future__ import annotations

from pathlib import Path

from openpyxl import load_workbook

TYPE_MATERIAL = "materialNumber"
TYPE_MFR = "MFRPartNo"
TYPE_UPC = "UPCCode"
TYPE_ALIAS = "Alias"
TYPE_UNKNOWN = "Unknown"

# Prefer product identity fields over secondary identifiers when a value
# appears in more than one column across the catalog.
_PRIORITY = {
    TYPE_MATERIAL: 1,
    TYPE_UPC: 2,
    TYPE_MFR: 3,
    TYPE_ALIAS: 4,
}


def _split_alias(value: str) -> list[str]:
    parts: list[str] = []
    for chunk in str(value or "").replace(";", ",").split(","):
        item = chunk.strip()
        if item:
            parts.append(item)
    return parts


def build_identifier_lookup(product_dataset_path: Path) -> dict[str, str]:
    """Map identifier value -> type using product_dataset.xlsx."""
    lookup: dict[str, str] = {}
    if not product_dataset_path.is_file():
        return lookup

    workbook = load_workbook(product_dataset_path, read_only=True, data_only=True)
    try:
        sheet = (
            workbook["Products"]
            if "Products" in workbook.sheetnames
            else workbook.active
        )
        rows = sheet.iter_rows(values_only=True)
        header = next(rows, None)
        if not header:
            return lookup
        headers = [str(cell or "").strip() for cell in header]
        idx = {name: i for i, name in enumerate(headers)}

        def _cell(row: tuple, name: str) -> str:
            if name not in idx or idx[name] >= len(row):
                return ""
            return "" if row[idx[name]] is None else str(row[idx[name]]).strip()

        def _remember(value: str, ident_type: str) -> None:
            if not value:
                return
            current = lookup.get(value)
            if current is None or _PRIORITY[ident_type] < _PRIORITY.get(current, 99):
                lookup[value] = ident_type

        for row in rows:
            if not row:
                continue
            _remember(_cell(row, "materialNumber"), TYPE_MATERIAL)
            _remember(_cell(row, "UPCCode"), TYPE_UPC)
            _remember(_cell(row, "MFRPartNo"), TYPE_MFR)
            for alias in _split_alias(_cell(row, "Alias")):
                _remember(alias, TYPE_ALIAS)
    finally:
        workbook.close()

    return lookup


def classify_identifier(value: str, lookup: dict[str, str]) -> str:
    text = (value or "").strip()
    if not text:
        return TYPE_UNKNOWN
    # Provided SKUs are treated as materialNumber.
    return lookup.get(text, TYPE_MATERIAL)
